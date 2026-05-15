"""
gerar_relatorio_ppl.py
Gera um arquivo Excel com tabelas comparando os resultados de perplexidade (PPL)
dos experimentos. Suporta múltiplos modelos e ações automaticamente.

Padrão dos arquivos: <Modelo>-<Ação>.json
Uso: python gerar_relatorio_ppl.py [pasta_com_jsons] [arquivo_saida.xlsx]
"""

import json
import sys
import glob
import os
from pathlib import Path
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Configurações de estilo ──────────────────────────────────────────────────

FONT_NAME = "Arial"

COLOR_HEADER_MODEL   = "1F4E79"  # azul escuro — cabeçalho de modelo
COLOR_HEADER_LANG    = "2E75B6"  # azul médio  — cabeçalho de língua
COLOR_HEADER_ACTION  = "D6E4F0"  # azul claro  — cabeçalho de ação
COLOR_BEST           = "C6EFCE"  # verde        — menor PPL (melhor)
COLOR_WORST          = "FFC7CE"  # vermelho     — maior PPL (pior)
COLOR_SUBHEADER      = "BDD7EE"  # azul pálido  — sub-cabeçalhos
COLOR_WHITE          = "FFFFFF"
COLOR_ALT_ROW        = "F2F7FC"  # linhas alternadas

thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def cell_style(ws, row, col, value=None, bold=False, font_color="000000",
               bg_color=None, align="center", wrap=False, num_fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=FONT_NAME, bold=bold, color=font_color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = BORDER
    if bg_color:
        cell.fill = PatternFill("solid", start_color=bg_color)
    if num_fmt:
        cell.number_format = num_fmt
    return cell


# ── Carregamento dos dados ───────────────────────────────────────────────────

def load_experiments(folder: str) -> dict:
    """
    Retorna estrutura: data[modelo][acao][lingua] = {"avg_ppl": float, "texts": [...]}
    """
    data = defaultdict(lambda: defaultdict(dict))
    pattern = os.path.join(folder, "*.json")
    files = glob.glob(pattern)
    if not files:
        raise FileNotFoundError(f"Nenhum arquivo .json encontrado em: {folder}")

    for fpath in files:
        stem = Path(fpath).stem          # ex: "Gemma-FAZER"
        parts = stem.split("-", 1)
        if len(parts) != 2:
            print(f"  [aviso] Ignorando arquivo com nome inesperado: {fpath}")
            continue
        modelo, acao = parts[0], parts[1]
        with open(fpath, encoding="utf-8") as f:
            content = json.load(f)
        for lingua, vals in content.items():
            data[modelo][acao][lingua] = vals

    return data


# ── Aba 1: Visão Geral — avg_ppl por modelo × língua × ação ─────────────────

def sheet_visao_geral(wb, data, modelos, acoes, linguas):
    ws = wb.create_sheet("Visão Geral")
    ws.freeze_panes = "C3"

    row = 1
    # Título
    ws.merge_cells(start_row=row, start_column=1,
                   end_row=row, end_column=2 + len(acoes))
    c = ws.cell(row=row, column=1,
                value="Perplexidade Média (avg_ppl) por Modelo, Língua e Ação")
    c.font = Font(name=FONT_NAME, bold=True, size=13, color=COLOR_WHITE)
    c.fill = PatternFill("solid", start_color=COLOR_HEADER_MODEL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 22
    row += 1

    for modelo in modelos:
        # Cabeçalho do modelo
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=2 + len(acoes))
        c = ws.cell(row=row, column=1, value=f"Modelo: {modelo}")
        c.font = Font(name=FONT_NAME, bold=True, color=COLOR_WHITE)
        c.fill = PatternFill("solid", start_color=COLOR_HEADER_LANG)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18
        row += 1

        # Cabeçalhos de coluna
        cell_style(ws, row, 1, "Língua", bold=True, bg_color=COLOR_SUBHEADER)
        cell_style(ws, row, 2, "Média Geral", bold=True, bg_color=COLOR_SUBHEADER)
        for j, acao in enumerate(acoes, start=3):
            cell_style(ws, row, j, acao, bold=True, bg_color=COLOR_SUBHEADER)
        row += 1

        header_row = row - 1

        # Coleta todos os valores da tabela para colorir melhor/pior
        tabela_vals = {}
        for lingua in linguas:
            vals_por_acao = []
            for acao in acoes:
                v = (data.get(modelo, {})
                         .get(acao, {})
                         .get(lingua, {})
                         .get("avg_ppl"))
                vals_por_acao.append(v)
            tabela_vals[lingua] = vals_por_acao

        all_valid = [v for row_vals in tabela_vals.values()
                     for v in row_vals if v is not None]
        global_min = min(all_valid) if all_valid else None
        global_max = max(all_valid) if all_valid else None

        for i, lingua in enumerate(linguas):
            bg = COLOR_ALT_ROW if i % 2 == 0 else COLOR_WHITE
            cell_style(ws, row, 1, lingua, bold=True, align="left", bg_color=bg)

            vals = tabela_vals[lingua]
            valid = [v for v in vals if v is not None]

            # Fórmula de média na coluna 2
            col_letters = [get_column_letter(3 + j) for j, v in enumerate(vals) if v is not None]
            if col_letters:
                refs = ",".join(f"{cl}{row}" for cl in col_letters)
                ws.cell(row=row, column=2).value = f"=AVERAGE({refs})"
                ws.cell(row=row, column=2).number_format = "0.00"
                ws.cell(row=row, column=2).font = Font(name=FONT_NAME)
                ws.cell(row=row, column=2).fill = PatternFill("solid", start_color=bg)
                ws.cell(row=row, column=2).border = BORDER
                ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")
            else:
                cell_style(ws, row, 2, "N/A", bg_color=bg)

            for j, v in enumerate(vals, start=3):
                if v is None:
                    cell_style(ws, row, j, "N/A", bg_color=bg)
                else:
                    if v == global_min:
                        bg_c = COLOR_BEST
                    elif v == global_max:
                        bg_c = COLOR_WORST
                    else:
                        bg_c = bg
                    cell_style(ws, row, j, v, bg_color=bg_c, num_fmt="0.00")
            row += 1

        # Linha de média por ação (rodapé)
        cell_style(ws, row, 1, "Média por Ação", bold=True, bg_color=COLOR_SUBHEADER)
        # média geral de todas as ações
        refs_media = ",".join(
            f"{get_column_letter(3+j)}{row-len(linguas)}:{get_column_letter(3+j)}{row-1}"
            for j in range(len(acoes))
        )
        ws.cell(row=row, column=2).value = f"=AVERAGE({refs_media})"
        ws.cell(row=row, column=2).number_format = "0.00"
        ws.cell(row=row, column=2).font = Font(name=FONT_NAME, bold=True)
        ws.cell(row=row, column=2).fill = PatternFill("solid", start_color=COLOR_SUBHEADER)
        ws.cell(row=row, column=2).border = BORDER
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")

        for j in range(len(acoes)):
            col = 3 + j
            r_start = row - len(linguas)
            r_end   = row - 1
            col_l   = get_column_letter(col)
            ws.cell(row=row, column=col).value = f"=AVERAGE({col_l}{r_start}:{col_l}{r_end})"
            ws.cell(row=row, column=col).number_format = "0.00"
            ws.cell(row=row, column=col).font = Font(name=FONT_NAME, bold=True)
            ws.cell(row=row, column=col).fill = PatternFill("solid", start_color=COLOR_SUBHEADER)
            ws.cell(row=row, column=col).border = BORDER
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")
        row += 2  # espaço entre modelos

    # Larguras
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14
    for j in range(len(acoes)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 16


# ── Aba 2: Comparação entre modelos por ação ────────────────────────────────

def sheet_por_acao(wb, data, modelos, acoes, linguas):
    for acao in acoes:
        ws = wb.create_sheet(f"Ação_{acao[:20]}")
        ws.freeze_panes = "B3"

        row = 1
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=1 + len(modelos))
        c = ws.cell(row=row, column=1,
                    value=f"Ação: {acao} — avg_ppl por Língua e Modelo")
        c.font = Font(name=FONT_NAME, bold=True, size=12, color=COLOR_WHITE)
        c.fill = PatternFill("solid", start_color=COLOR_HEADER_MODEL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1

        # Cabeçalhos
        cell_style(ws, row, 1, "Língua", bold=True, bg_color=COLOR_SUBHEADER)
        for j, modelo in enumerate(modelos, start=2):
            cell_style(ws, row, j, modelo, bold=True, bg_color=COLOR_SUBHEADER)
        row += 1

        # Coleta valores para min/max global
        all_vals = []
        for lingua in linguas:
            for modelo in modelos:
                v = data.get(modelo, {}).get(acao, {}).get(lingua, {}).get("avg_ppl")
                if v is not None:
                    all_vals.append(v)
        g_min = min(all_vals) if all_vals else None
        g_max = max(all_vals) if all_vals else None

        for i, lingua in enumerate(linguas):
            bg = COLOR_ALT_ROW if i % 2 == 0 else COLOR_WHITE
            cell_style(ws, row, 1, lingua, bold=True, align="left", bg_color=bg)
            for j, modelo in enumerate(modelos, start=2):
                v = data.get(modelo, {}).get(acao, {}).get(lingua, {}).get("avg_ppl")
                if v is None:
                    cell_style(ws, row, j, "N/A", bg_color=bg)
                else:
                    if v == g_min:
                        bg_c = COLOR_BEST
                    elif v == g_max:
                        bg_c = COLOR_WORST
                    else:
                        bg_c = bg
                    cell_style(ws, row, j, v, bg_color=bg_c, num_fmt="0.00")
            row += 1

        # Rodapé — média por modelo
        cell_style(ws, row, 1, "Média", bold=True, bg_color=COLOR_SUBHEADER)
        for j in range(len(modelos)):
            col = 2 + j
            r_start = 3
            r_end   = row - 1
            col_l   = get_column_letter(col)
            c = ws.cell(row=row, column=col)
            c.value = f"=AVERAGE({col_l}{r_start}:{col_l}{r_end})"
            c.number_format = "0.00"
            c.font = Font(name=FONT_NAME, bold=True)
            c.fill = PatternFill("solid", start_color=COLOR_SUBHEADER)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")

        ws.column_dimensions["A"].width = 18
        for j in range(len(modelos)):
            ws.column_dimensions[get_column_letter(2 + j)].width = 20


# ── Aba 3: Detalhamento por texto ────────────────────────────────────────────

def sheet_detalhamento(wb, data, modelos, acoes, linguas):
    ws = wb.create_sheet("Detalhamento por Texto")
    ws.freeze_panes = "E2"

    headers = ["Modelo", "Ação", "Língua", "text_id", "num_tokens", "ppl"]
    for j, h in enumerate(headers, start=1):
        cell_style(ws, 1, j, h, bold=True, bg_color=COLOR_HEADER_LANG,
                   font_color=COLOR_WHITE)

    row = 2
    for modelo in modelos:
        for acao in acoes:
            for lingua in linguas:
                entry = data.get(modelo, {}).get(acao, {}).get(lingua, {})
                texts = entry.get("texts", [])
                for t in texts:
                    bg = COLOR_ALT_ROW if row % 2 == 0 else COLOR_WHITE
                    cell_style(ws, row, 1, modelo, align="left", bg_color=bg)
                    cell_style(ws, row, 2, acao, align="left", bg_color=bg)
                    cell_style(ws, row, 3, lingua, align="left", bg_color=bg)
                    cell_style(ws, row, 4, t.get("text_id"), bg_color=bg)
                    cell_style(ws, row, 5, t.get("num_tokens"), bg_color=bg)
                    cell_style(ws, row, 6, t.get("ppl"), bg_color=bg, num_fmt="0.00")
                    row += 1

    widths = [14, 18, 14, 10, 12, 12]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


# ── Aba 4: Ranking de modelos por língua ─────────────────────────────────────

def sheet_ranking(wb, data, modelos, acoes, linguas):
    ws = wb.create_sheet("Ranking por Língua")
    ws.freeze_panes = "A3"

    row = 1
    ws.merge_cells(start_row=row, start_column=1,
                   end_row=row, end_column=3)
    c = ws.cell(row=row, column=1,
                value="Ranking de Modelos por Língua (média de avg_ppl sobre todas as ações)")
    c.font = Font(name=FONT_NAME, bold=True, size=12, color=COLOR_WHITE)
    c.fill = PatternFill("solid", start_color=COLOR_HEADER_MODEL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 20
    row += 1

    cell_style(ws, row, 1, "Língua", bold=True, bg_color=COLOR_SUBHEADER)
    cell_style(ws, row, 2, "Ranking", bold=True, bg_color=COLOR_SUBHEADER)
    cell_style(ws, row, 3, "Modelo (avg_ppl médio)", bold=True,
               bg_color=COLOR_SUBHEADER)
    row += 1

    for lingua in linguas:
        medias = {}
        for modelo in modelos:
            vals = []
            for acao in acoes:
                v = data.get(modelo, {}).get(acao, {}).get(lingua, {}).get("avg_ppl")
                if v is not None:
                    vals.append(v)
            if vals:
                medias[modelo] = sum(vals) / len(vals)

        ranking = sorted(medias.items(), key=lambda x: x[1])
        start_row = row
        for rank, (modelo, media) in enumerate(ranking, start=1):
            bg = COLOR_BEST if rank == 1 else (COLOR_WORST if rank == len(ranking)
                                               else (COLOR_ALT_ROW if rank % 2 == 0
                                                     else COLOR_WHITE))
            cell_style(ws, row, 1, lingua if rank == 1 else "", bold=(rank == 1),
                       align="left", bg_color=bg)
            cell_style(ws, row, 2, rank, bg_color=bg)
            cell_style(ws, row, 3, f"{modelo}  ({media:.2f})", align="left",
                       bg_color=bg)
            row += 1

        if len(ranking) > 1:
            ws.merge_cells(start_row=start_row, start_column=1,
                           end_row=row - 1, end_column=1)
            ws.cell(row=start_row, column=1).alignment = Alignment(
                horizontal="left", vertical="center")
        row += 1  # linha em branco

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 32


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else "./casualLM"
    output = sys.argv[2] if len(sys.argv) > 2 else "./relatorio_ppl_casualLM.xlsx"

    print(f"Carregando JSONs de: {folder}")
    data = load_experiments(folder)

    modelos = sorted(data.keys())
    acoes   = sorted({a for m in data.values() for a in m.keys()})
    linguas = sorted({l for m in data.values() for a in m.values() for l in a.keys()})

    print(f"  Modelos : {modelos}")
    print(f"  Ações   : {acoes}")
    print(f"  Línguas : {linguas}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba padrão

    print("Gerando abas...")
    sheet_visao_geral(wb, data, modelos, acoes, linguas)
    sheet_por_acao(wb, data, modelos, acoes, linguas)
    sheet_detalhamento(wb, data, modelos, acoes, linguas)
    sheet_ranking(wb, data, modelos, acoes, linguas)

    os.makedirs(os.path.dirname(output), exist_ok=True)
    wb.save(output)
    print(f"Arquivo salvo em: {output}")


if __name__ == "__main__":
    main()
