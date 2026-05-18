#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Gera uma planilha Excel comparando resultados de perplexidade a partir de vários JSONs.

Nesta versão, a aba Comparacao_AvgPPL compara cada língua com o PB NATIVO
da mesma tarefa/verbo. Assim, cada grupo de tarefa tem sua própria linha-base.

Formato esperado dos JSONs:
{
  "TAREFA": {
    "LÍNGUA": {
      "avg_ppl": 123.45,
      "texts": [
        {"text_id": 0, "num_tokens": 100, "ppl": 12.3}
      ]
    }
  }
}

Uso:
    python gerar_excel_modelos.py --input ./resultados_json --output comparacao_modelos.xlsx

Dependências:
    pip install pandas openpyxl
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference


BASELINE_LINGUA = "PB NATIVO"


def carregar_jsons(pasta: Path) -> Dict[str, Dict[str, Any]]:
    """Carrega todos os JSONs da pasta. O nome do modelo vem do nome do arquivo."""
    if not pasta.exists():
        raise FileNotFoundError(f"Pasta não encontrada: {pasta}")

    arquivos = sorted(pasta.glob("*.json"))
    if not arquivos:
        raise FileNotFoundError(f"Nenhum arquivo .json encontrado em: {pasta}")

    modelos: Dict[str, Dict[str, Any]] = {}

    for arquivo in arquivos:
        nome_modelo = arquivo.stem
        with arquivo.open("r", encoding="utf-8") as f:
            try:
                modelos[nome_modelo] = json.load(f)
            except json.JSONDecodeError as e:
                raise ValueError(f"Erro ao ler JSON em {arquivo.name}: {e}") from e

    return modelos


def iter_metricas_agregadas(modelos: Dict[str, Dict[str, Any]]) -> Iterable[dict]:
    """Transforma os avg_ppl dos JSONs em linhas tabulares."""
    for modelo, dados in modelos.items():
        for tarefa, linguas in dados.items():
            if not isinstance(linguas, dict):
                continue

            for lingua, info in linguas.items():
                if not isinstance(info, dict):
                    continue

                textos = info.get("texts", [])

                yield {
                    "tarefa": tarefa,
                    "lingua": lingua,
                    "modelo": modelo,
                    "avg_ppl": info.get("avg_ppl"),
                    "qtd_textos": len(textos) if isinstance(textos, list) else 0,
                }


def iter_metricas_textos(modelos: Dict[str, Dict[str, Any]]) -> Iterable[dict]:
    """Transforma as métricas por texto em linhas tabulares."""
    for modelo, dados in modelos.items():
        for tarefa, linguas in dados.items():
            if not isinstance(linguas, dict):
                continue

            for lingua, info in linguas.items():
                if not isinstance(info, dict):
                    continue

                textos = info.get("texts", [])
                if not isinstance(textos, list):
                    continue

                for texto in textos:
                    yield {
                        "tarefa": tarefa,
                        "lingua": lingua,
                        "text_id": texto.get("text_id"),
                        "modelo": modelo,
                        "num_tokens": texto.get("num_tokens"),
                        "ppl": texto.get("ppl"),
                    }


def criar_comparacao_relativa_pb(df_avg: pd.DataFrame) -> pd.DataFrame:
    """
    Cria a tabela principal de comparação.

    Para cada modelo e cada tarefa/verbo:
    - localiza o avg_ppl do PB NATIVO;
    - calcula a diferença percentual de cada língua em relação a esse PB NATIVO.

    Fórmula:
        vs_PB_% = ((avg_ppl_lingua / avg_ppl_pb_nativo_da_tarefa) - 1) * 100

    Interpretação:
        valor negativo -> menor PPL que PB NATIVO
        valor zero     -> igual ao PB NATIVO
        valor positivo -> maior PPL que PB NATIVO
    """
    modelos_ordenados = sorted(df_avg["modelo"].dropna().unique())

    comp_avg = (
        df_avg
        .pivot_table(
            index=["tarefa", "lingua"],
            columns="modelo",
            values="avg_ppl",
            aggfunc="first",
        )
        .reset_index()
    )

    # Tabela com o PB NATIVO de cada tarefa para cada modelo.
    baseline = (
        df_avg[df_avg["lingua"] == BASELINE_LINGUA]
        .pivot_table(
            index="tarefa",
            columns="modelo",
            values="avg_ppl",
            aggfunc="first",
        )
    )

    # Cria colunas intercaladas: Modelo_avg_ppl, Modelo_vs_PB_%.
    resultado = comp_avg[["tarefa", "lingua"]].copy()

    for modelo in modelos_ordenados:
        if modelo not in comp_avg.columns:
            continue

        resultado[f"{modelo}_avg_ppl"] = comp_avg[modelo]

        def calcular_vs_pb(row: pd.Series) -> float | None:
            tarefa = row["tarefa"]
            valor = row[modelo]

            if tarefa not in baseline.index:
                return None
            if modelo not in baseline.columns:
                return None

            valor_pb = baseline.loc[tarefa, modelo]

            if pd.isna(valor) or pd.isna(valor_pb) or valor_pb == 0:
                return None

            return ((valor / valor_pb) - 1) * 100

        resultado[f"{modelo}_vs_PB_%"] = comp_avg.apply(calcular_vs_pb, axis=1)

    # Ordena para manter PB NATIVO próximo dentro de cada tarefa.
    # PB NATIVO aparece primeiro no grupo para servir como baseline visual.
    resultado["_ordem_lingua"] = resultado["lingua"].apply(lambda x: 0 if x == BASELINE_LINGUA else 1)
    resultado = (
        resultado
        .sort_values(["tarefa", "_ordem_lingua", "lingua"])
        .drop(columns=["_ordem_lingua"])
        .reset_index(drop=True)
    )

    return resultado


def criar_ranking(df_avg: pd.DataFrame) -> pd.DataFrame:
    """Cria ranking do melhor modelo por tarefa e língua, considerando menor avg_ppl."""
    comp_avg = (
        df_avg
        .pivot_table(
            index=["tarefa", "lingua"],
            columns="modelo",
            values="avg_ppl",
            aggfunc="first",
        )
        .reset_index()
    )

    modelos_ordenados = sorted([c for c in comp_avg.columns if c not in ["tarefa", "lingua"]])

    linhas_ranking: List[dict] = []
    for _, row in comp_avg.iterrows():
        valores = {
            modelo: row[modelo]
            for modelo in modelos_ordenados
            if pd.notna(row[modelo])
        }

        if not valores:
            continue

        ordenados = sorted(valores.items(), key=lambda x: x[1])
        melhor_modelo, melhor_ppl = ordenados[0]
        segundo_ppl = ordenados[1][1] if len(ordenados) > 1 else None

        if segundo_ppl is not None and pd.notna(melhor_ppl) and segundo_ppl != 0:
            vantagem_pct = ((segundo_ppl - melhor_ppl) / segundo_ppl) * 100
        else:
            vantagem_pct = None

        linhas_ranking.append({
            "tarefa": row["tarefa"],
            "lingua": row["lingua"],
            "melhor_modelo": melhor_modelo,
            "menor_avg_ppl": melhor_ppl,
            "segundo_menor_avg_ppl": segundo_ppl,
            "vantagem_%_vs_segundo": vantagem_pct,
        })

    return pd.DataFrame(linhas_ranking)


def criar_detalhes_textos(df_textos: pd.DataFrame) -> pd.DataFrame:
    """Cria tabela detalhada por texto, intercalando PPL e tokens por modelo."""
    if df_textos.empty:
        return pd.DataFrame()

    modelos_ordenados = sorted(df_textos["modelo"].dropna().unique())

    ppl_wide = (
        df_textos
        .pivot_table(
            index=["tarefa", "lingua", "text_id"],
            columns="modelo",
            values="ppl",
            aggfunc="first",
        )
    )

    tokens_wide = (
        df_textos
        .pivot_table(
            index=["tarefa", "lingua", "text_id"],
            columns="modelo",
            values="num_tokens",
            aggfunc="first",
        )
    )

    detalhes = pd.DataFrame(index=ppl_wide.index)

    for modelo in modelos_ordenados:
        if modelo in ppl_wide.columns:
            detalhes[f"{modelo}_ppl"] = ppl_wide[modelo]
        if modelo in tokens_wide.columns:
            detalhes[f"{modelo}_tokens"] = tokens_wide[modelo]

    return detalhes.reset_index()


def criar_tabelas(modelos: Dict[str, Dict[str, Any]]) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Cria DataFrames: base agregada, comparação relativa ao PB, ranking e detalhes."""
    df_avg = pd.DataFrame(iter_metricas_agregadas(modelos))
    df_textos = pd.DataFrame(iter_metricas_textos(modelos))

    if df_avg.empty:
        raise ValueError("Nenhuma métrica agregada encontrada nos JSONs.")

    comp_avg = criar_comparacao_relativa_pb(df_avg)
    ranking = criar_ranking(df_avg)
    detalhes = criar_detalhes_textos(df_textos)

    return df_avg, comp_avg, ranking, detalhes


def ajustar_visual_excel(caminho_excel: Path) -> None:
    """Aplica formatação para deixar a planilha fácil de comparar."""
    wb = load_workbook(caminho_excel)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    baseline_fill = PatternFill("solid", fgColor="D9EAF7")
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        for row in ws.iter_rows(min_row=2):
            is_pb_baseline = ws.cell(row=row[0].row, column=2).value == BASELINE_LINGUA

            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center")

                if is_pb_baseline:
                    cell.fill = baseline_fill

                if isinstance(cell.value, float):
                    if isinstance(ws.cell(row=1, column=cell.column).value, str) and ws.cell(row=1, column=cell.column).value.endswith("_vs_PB_%"):
                        cell.number_format = '0.00"%"'
                    else:
                        cell.number_format = '#,##0.00'

        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))

            width = min(max(max_len + 2, 12), 34)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        if ws.max_row >= 2 and ws.max_column >= 2:
            table_name = ("tbl_" + "".join(ch if ch.isalnum() else "_" for ch in ws.title))[:31]
            ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

            tab = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            tab.tableStyleInfo = style
            ws.add_table(tab)

    # Na comparação principal:
    # - colunas *_avg_ppl recebem escala normal: menor verde, maior vermelho;
    # - colunas *_vs_PB_% recebem escala relativa: negativo verde, zero amarelo, positivo vermelho.
    if "Comparacao_AvgPPL" in wb.sheetnames:
        ws = wb["Comparacao_AvgPPL"]

        for col_idx in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col_idx).value
            if not isinstance(header, str) or ws.max_row < 2:
                continue

            col_letter = get_column_letter(col_idx)
            col_range = f"{col_letter}2:{col_letter}{ws.max_row}"

            if header.endswith("_avg_ppl"):
                ws.conditional_formatting.add(
                    col_range,
                    ColorScaleRule(
                        start_type="min", start_color="63BE7B",
                        mid_type="percentile", mid_value=50, mid_color="FFEB84",
                        end_type="max", end_color="F8696B",
                    ),
                )

            elif header.endswith("_vs_PB_%"):
                # Como a coluna é diferença percentual relativa ao PB NATIVO:
                # verde = menor que PB NATIVO
                # amarelo = igual ao PB NATIVO
                # vermelho = maior que PB NATIVO
                ws.conditional_formatting.add(
                    col_range,
                    ColorScaleRule(
                        start_type="min", start_color="63BE7B",
                        mid_type="num", mid_value=0, mid_color="FFEB84",
                        end_type="max", end_color="F8696B",
                    ),
                )

    # Em detalhes, colore apenas colunas de PPL.
    if "Detalhes_Textos" in wb.sheetnames:
        ws = wb["Detalhes_Textos"]

        for col_idx in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col_idx).value
            if isinstance(header, str) and header.endswith("_ppl") and ws.max_row > 1:
                col_letter = get_column_letter(col_idx)
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    ColorScaleRule(
                        start_type="min", start_color="63BE7B",
                        mid_type="percentile", mid_value=50, mid_color="FFEB84",
                        end_type="max", end_color="F8696B",
                    ),
                )

    # Gráfico simples na visão geral.
    if "Visao_Geral" in wb.sheetnames:
        ws = wb["Visao_Geral"]
        if ws.max_row >= 2:
            chart = BarChart()
            chart.title = "Avg PPL geral por modelo"
            chart.y_axis.title = "Avg PPL"
            chart.x_axis.title = "Modelo"

            data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 8
            chart.width = 16
            ws.add_chart(chart, "D2")

    wb.save(caminho_excel)


def gerar_excel(pasta_jsons: Path, saida_excel: Path) -> None:
    modelos = carregar_jsons(pasta_jsons)
    df_avg, comp_avg, ranking, detalhes = criar_tabelas(modelos)

    visao_geral = (
        df_avg
        .groupby("modelo", as_index=False)
        .agg(
            avg_ppl_geral=("avg_ppl", "mean"),
            mediana_avg_ppl=("avg_ppl", "median"),
            qtd_combinacoes=("avg_ppl", "count"),
        )
        .sort_values("avg_ppl_geral", ascending=True)
    )

    saida_excel.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(saida_excel, engine="openpyxl") as writer:
        visao_geral.to_excel(writer, sheet_name="Visao_Geral", index=False)
        comp_avg.to_excel(writer, sheet_name="Comparacao_AvgPPL", index=False)
        ranking.to_excel(writer, sheet_name="Ranking", index=False)
        detalhes.to_excel(writer, sheet_name="Detalhes_Textos", index=False)
        df_avg.to_excel(writer, sheet_name="Base_Agregada", index=False)

    ajustar_visual_excel(saida_excel)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Gera uma planilha Excel comparando métricas de vários modelos a partir de JSONs."
    )
    parser.add_argument(
        "--input",
        "-i",
        required=True,
        help="Pasta contendo os arquivos .json dos modelos.",
    )
    parser.add_argument(
        "--output",
        "-o",
        default="comparacao_modelos.xlsx",
        help="Caminho do arquivo Excel de saída.",
    )

    args = parser.parse_args()

    pasta_jsons = Path(args.input)
    saida_excel = Path(args.output)

    gerar_excel(pasta_jsons, saida_excel)
    print(f"Excel gerado com sucesso: {saida_excel.resolve()}")


if __name__ == "__main__":
    main()
